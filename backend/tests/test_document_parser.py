"""Document parser tests — verify content preservation across formats."""

import os
import tempfile
import sys

# Ensure the backend root is on sys.path for imports
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.parser.factory import ParserFactory


def _make_temp_file(content: str, suffix: str) -> str:
    """Create a temp file with *content* and *.suffix*, return its path."""
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ── TXT Parser Tests ──────────────────────────────────────────────


class TestTextParser:
    """Verify that .txt parsing preserves digits, times, dates, whitespace."""

    def test_simple_text(self):
        content = "Hello World"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert result["content"] == "Hello World"
        assert result["metadata"]["type"] == "txt"

    def test_time_range_preserved(self):
        """08:00-12:00 must survive parsing untouched."""
        content = "门诊时间：08:00-12:00\n下午 14:00-17:30"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "08:00-12:00" in result["content"]
        assert "14:00-17:30" in result["content"]

    def test_date_preserved(self):
        content = "日期：2026-07-19"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "2026-07-19" in result["content"]

    def test_currency_preserved(self):
        content = "金额：100元"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "100元" in result["content"]

    def test_percentage_preserved(self):
        content = "完成度：50%"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "50%" in result["content"]

    def test_digit_count(self):
        """Parser output must not lose digits."""
        content = "08:00-12:00, 100元, 50%, 2026-07-19"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        original_digits = sum(c.isdigit() for c in content)
        result_digits = sum(c.isdigit() for c in result["content"])
        assert result_digits == original_digits, (
            f"Digit loss: {original_digits} → {result_digits}"
        )

    def test_whitespace_preserved(self):
        """Newlines and spaces must not be collapsed."""
        content = "line1\n\nline2\n  indented"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "line1\n\nline2" in result["content"]

    def test_chinese_time_range(self):
        content = "8点到12点"
        path = _make_temp_file(content, ".txt")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "8点到12点" in result["content"]


# ── Markdown Parser Tests ─────────────────────────────────────────


class TestMarkdownParser:
    def test_basic_markdown(self):
        content = "# Title\n\nSome **bold** text"
        path = _make_temp_file(content, ".md")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "# Title" in result["content"]
        assert "**bold**" in result["content"]
        assert result["metadata"]["type"] == "md"

    def test_table_not_destroyed(self):
        content = "| A | B |\n|---|---|\n| 1 | 2 |"
        path = _make_temp_file(content, ".md")
        result = ParserFactory.get_parser(path).parse(path)
        os.remove(path)
        assert "|" in result["content"]
        assert "1" in result["content"]


# ── DOCX Parser Tests (if python-docx available) ──────────────────


class TestDocxParser:
    """Test .docx parsing.  Requires python-docx."""

    def _create_simple_docx(self) -> str:
        """Create a .docx with a paragraph and a table, return its path."""
        try:
            from docx import Document
            from docx.shared import Inches
        except ImportError:
            return None  # skip

        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)

        doc = Document()
        doc.add_paragraph("门诊时间：")
        doc.add_paragraph("上午 08:00-12:00")
        doc.add_paragraph("下午 14:00-17:30")

        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "时间"
        table.cell(0, 1).text = "上午"
        table.cell(0, 2).text = "下午"
        table.cell(1, 0).text = "周一"
        table.cell(1, 1).text = "08:00"
        table.cell(1, 2).text = "14:00"

        doc.save(path)
        return path

    def test_paragraphs_and_tables(self):
        path = self._create_simple_docx()
        if path is None:
            return  # python-docx not installed, skip
        try:
            result = ParserFactory.get_parser(path).parse(path)
            content = result["content"]
            assert "门诊时间" in content
            assert "08:00-12:00" in content
            assert "14:00-17:30" in content
            assert "时间 | 上午 | 下午" in content, f"Table header missing in:\n{content}"
            assert "周一 | 08:00 | 14:00" in content, f"Table row missing in:\n{content}"
            assert result["metadata"]["tables"] >= 1
            # Verify digit count preserved
            digit_count = sum(c.isdigit() for c in content)
            assert digit_count >= 12, f"Digit loss: only {digit_count} digits in:\n{content}"
        finally:
            os.remove(path)


# ── DOC Parser Tests (legacy .doc, via olefile) ────────────────────


def _make_minimal_doc(text: str, is_unicode: bool = True) -> bytes:
    """Build a minimal WordDocument stream (FIB header + body)."""
    import struct

    body = text.encode("utf-16-le") if is_unicode else text.encode("gb18030")
    fc_min = 0x20
    fc_mac = fc_min + len(body)

    fib = bytearray(32)
    struct.pack_into("<H", fib, 0x00, 0xA5EC)  # wIdent
    struct.pack_into("<H", fib, 0x02, 0x00C1)  # nFib (Word 97)
    struct.pack_into("<H", fib, 0x0A, 0x1000 if is_unicode else 0x0000)  # flags
    struct.pack_into("<I", fib, 0x18, fc_min)
    struct.pack_into("<I", fib, 0x1C, fc_mac)
    return bytes(fib) + body


class TestDocParser:
    """Test legacy .doc parsing via DocParser."""

    def test_extract_unicode_body(self):
        from app.services.parser.doc_parser import DocParser

        stream = _make_minimal_doc("门诊时间：08:00-12:00\r下午 14:00-17:30")
        content = DocParser()._extract_body(stream)
        assert "门诊时间" in content
        assert "08:00-12:00" in content
        assert "14:00-17:30" in content

    def test_extract_single_byte_body(self):
        from app.services.parser.doc_parser import DocParser

        stream = _make_minimal_doc("金额：100元\r完成度：50%", is_unicode=False)
        content = DocParser()._extract_body(stream)
        assert "100元" in content
        assert "50%" in content

    def test_clean_text_normalises_control_chars(self):
        from app.services.parser.doc_parser import DocParser

        cleaned = DocParser._clean_text("A\rB\x0bC\x0cD\x07E")
        assert "\r" not in cleaned
        assert "\x07" not in cleaned
        assert " | " in cleaned

    def test_decode_single_byte_prefers_gbk(self):
        from app.services.parser.doc_parser import DocParser

        decoded = DocParser._decode_single_byte("中文测试".encode("gbk"))
        assert "中文测试" in decoded

    def test_looks_like_text_filters_garbage(self):
        from app.services.parser.doc_parser import DocParser

        assert DocParser._looks_like_text("一、核心技术") is True
        assert DocParser._looks_like_text("\x00\x01") is False


# ── PDF Parser Tests (if PyMuPDF available) ───────────────────────


class TestPdfParser:
    """Test .pdf parsing.  Requires PyMuPDF (fitz)."""

    def _create_simple_pdf(self) -> str:
        import fitz  # PyMuPDF — declared in requirements.txt

        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Time: 08:00-12:00\nAfternoon: 14:00-17:30")
        doc.save(path)
        doc.close()
        return path

    def test_pdf_text_preserved(self):
        path = self._create_simple_pdf()
        try:
            result = ParserFactory.get_parser(path).parse(path)
            content = result["content"]
            assert "08:00-12:00" in content, f"Time range missing in:\n{content}"
            assert "14:00-17:30" in content, f"Time range missing in:\n{content}"
            assert result["metadata"]["pages"] >= 1
            digit_count = sum(c.isdigit() for c in content)
            assert digit_count >= 8, f"Digit loss: only {digit_count} digits in:\n{content}"
        finally:
            os.remove(path)


# ── XLSX Parser Tests ─────────────────────────────────────────────


class TestXlsxParser:
    """Test .xlsx parsing via openpyxl."""

    def test_rows_preserved(self):
        from openpyxl import Workbook

        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"
        ws.append(["时间", "上午", "下午"])
        ws.append(["周一", "08:00", "14:00"])
        wb.save(path)
        try:
            result = ParserFactory.get_parser(path).parse(path)
            content = result["content"]
            assert "08:00" in content
            assert "14:00" in content
            assert "时间 | 上午 | 下午" in content, f"Header missing in:\n{content}"
            assert result["metadata"]["type"] == "xlsx"
        finally:
            os.remove(path)


# ── CSV Parser Tests ──────────────────────────────────────────────


class TestCsvParser:
    """Test .csv parsing via the standard-library csv module."""

    def test_rows_preserved(self):
        content = "时间,上午,下午\n周一,08:00,14:00\n"
        path = _make_temp_file(content, ".csv")
        try:
            result = ParserFactory.get_parser(path).parse(path)
            parsed = result["content"]
            assert "08:00" in parsed
            assert "14:00" in parsed
            assert "时间 | 上午 | 下午" in parsed, f"Header missing in:\n{parsed}"
            assert result["metadata"]["type"] == "csv"
        finally:
            os.remove(path)


# ── Factory Tests ─────────────────────────────────────────────────


class TestParserFactory:
    def test_factory_returns_correct_parser(self):
        from app.services.parser.text_parser import TextParser
        from app.services.parser.markdown_parser import MarkdownParser
        from app.services.parser.docx_parser import DocxParser
        from app.services.parser.doc_parser import DocParser
        from app.services.parser.pdf_parser import PdfParser
        from app.services.parser.xlsx_parser import XlsxParser
        from app.services.parser.csv_parser import CsvParser

        assert isinstance(ParserFactory.get_parser("a.txt"), TextParser)
        assert isinstance(ParserFactory.get_parser("a.md"), MarkdownParser)
        assert isinstance(ParserFactory.get_parser("a.docx"), DocxParser)
        assert isinstance(ParserFactory.get_parser("a.doc"), DocParser)
        assert isinstance(ParserFactory.get_parser("a.pdf"), PdfParser)
        assert isinstance(ParserFactory.get_parser("a.xlsx"), XlsxParser)
        assert isinstance(ParserFactory.get_parser("a.csv"), CsvParser)

    def test_unsupported_extension_raises(self):
        import pytest
        with pytest.raises(ValueError, match="Unsupported file type"):
            ParserFactory.get_parser("a.xyz")