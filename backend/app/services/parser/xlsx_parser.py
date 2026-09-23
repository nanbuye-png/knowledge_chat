"""XLSX parser — extracts cell values from every worksheet.

Uses ``openpyxl`` to read all sheets row by row.  Cells within a row are
joined with `` | `` to preserve table structure for downstream chunking and
retrieval.  Fully-empty rows are dropped.
"""

from loguru import logger

from .base import DocumentParser


class XlsxParser(DocumentParser):
    """Parse .xlsx files, preserving sheet/row/column structure."""

    def parse(self, file_path: str) -> dict:
        try:
            import openpyxl
        except ImportError:
            raise ImportError(
                "openpyxl is required to parse .xlsx files. "
                "Install with: pip install openpyxl"
            )

        workbook = openpyxl.load_workbook(
            file_path, read_only=True, data_only=True
        )
        try:
            sheet_parts: list[str] = []
            table_count = 0
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows: list[str] = [f"[工作表: {sheet_name}]"]
                for row in sheet.iter_rows(values_only=True):
                    cells = [
                        "" if v is None else str(v).strip() for v in row
                    ]
                    if not any(cells):
                        continue
                    rows.append(" | ".join(cells))
                if len(rows) > 1:
                    table_count += 1
                sheet_parts.append("\n".join(rows))
        finally:
            workbook.close()

        content = "\n\n".join(sheet_parts)
        digit_count = sum(c.isdigit() for c in content)

        metadata = {
            "type": "xlsx",
            "pages": 0,
            "tables": table_count,
        }

        logger.debug("Document Parser Debug:")
        logger.debug(f"  filename: {file_path}")
        logger.debug(f"  type: xlsx")
        logger.debug(f"  table_count: {table_count}")
        logger.debug(f"  text_length: {len(content)}")
        logger.debug(f"  digit_count: {digit_count}")
        logger.debug(f"  preview: {repr(content[:200])}")

        return {
            "filename": file_path.split("\\")[-1].split("/")[-1],
            "content": content,
            "metadata": metadata,
        }
