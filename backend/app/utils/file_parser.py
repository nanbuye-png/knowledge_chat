import os
from loguru import logger
from ..core.config import settings


def parse_file(file_path: str) -> str:
    """
    Parse a document file and extract text content.
    Supports: PDF, DOCX, MD, TXT
    """
    ext = os.path.splitext(file_path)[1].lower()
    logger.info(f"Parsing file: {file_path} (type: {ext})")

    try:
        if ext == ".pdf":
            return _parse_pdf(file_path)
        elif ext in (".docx", ".doc"):
            return _parse_docx(file_path)
        elif ext == ".md":
            return _parse_markdown(file_path)
        elif ext == ".txt":
            return _parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    except Exception as e:
        logger.error(f"Failed to parse file {file_path}: {e}")
        raise


def _parse_pdf(file_path: str) -> str:
    """Parse PDF file using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("PyMuPDF is required for PDF parsing. Install with: pip install PyMuPDF")

    text_parts = []
    with fitz.open(file_path) as doc:
        for page_num, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                text_parts.append(f"[第{page_num + 1}页]\n{text.strip()}")

    full_text = "\n\n".join(text_parts)
    logger.info(f"PDF parsed: {len(text_parts)} pages, {len(full_text)} chars")
    return full_text


def _parse_docx(file_path: str) -> str:
    """Parse DOCX file using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")

    doc = Document(file_path)
    text_parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                text_parts.append(row_text)

    full_text = "\n\n".join(text_parts)
    logger.info(f"DOCX parsed: {len(text_parts)} paragraphs, {len(full_text)} chars")
    return full_text


def _parse_markdown(file_path: str) -> str:
    """Parse Markdown file."""
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    logger.info(f"Markdown parsed: {len(text)} chars")
    return text


def _parse_txt(file_path: str) -> str:
    """Parse plain text file."""
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    logger.info(f"TXT parsed: {len(text)} chars")
    return text